"""Danh sách HĐ lấy bằng nút "Xuất Excel" của cổng GDT, rồi mới đưa vào hộp chờ.

Kiểm: URL giống hệt cổng, đọc đúng cột (2 bố cục: HĐ thường / máy tính tiền), ngày lập đúng
múi giờ, GDT từ chối xuất file thì lấy lại qua API JSON, HĐ không mã không có XML thì nằm chờ
mail thay vì thử lại 10 ngày.

File thật của GDT (Apache POI) ghi <dimension ref="A1"/>. Mẫu test giả lập đúng điểm này, vì
openpyxl read_only tin con số đó và chỉ thấy 1 dòng → mất cả danh sách mà không báo lỗi.

Run::

    pytest tests/test_invoice_list_excel.py -q
"""

from __future__ import annotations

import importlib.util
import io
import re
import sys
import zipfile
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import openpyxl
import pytest
import requests

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import api_client           # noqa: E402
import invoice_list_excel   # noqa: E402
import job_policy           # noqa: E402


def _load_invoice_job():
    spec = importlib.util.spec_from_file_location("invoice_job", ROOT / "web" / "invoice_job.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


invoice_job = _load_invoice_job()

# ---------------------------------------------------------------------------
# File mẫu: cùng bố cục file GDT xuất 10/2026 (dữ liệu giả)
# ---------------------------------------------------------------------------

QUERY_HEADER = [
    "STT", "Ký hiệu mẫu số", "Ký hiệu hóa đơn", "Số hóa đơn", "Ngày lập",
    "MST người bán/MST người xuất hàng", "Tên người bán/Tên người xuất hàng",
    "MST người mua/MST người nhận hàng", "Tên người mua/Tên người nhận hàng", "Địa chỉ người mua",
    "Tổng tiền chưa thuế", "Tổng tiền thuế", "Tổng tiền chiết khấu thương mại", "Tổng tiền phí",
    "Tổng tiền thanh toán", "Đơn vị tiền tệ", "Tỷ giá", "Trạng thái hóa đơn",
    "Kết quả kiểm tra hóa đơn"]
SCO_HEADER = [
    "STT", "Ký hiệu mẫu số", "Ký hiệu hóa đơn", "Số hóa đơn", "Ngày lập",
    "MST người bán/MST người xuất hàng", "Tên người bán/Tên người xuất hàng", "Địa chỉ người bán",
    "MST người mua/MST người nhận hàng", "Tên người mua/Tên người nhận hàng", "Căn cước công dân",
    "Tổng tiền chưa thuế", "Tổng tiền thuế", "Tổng tiền chiết khấu thương mại",
    "Tổng tiền thanh toán", "Trạng thái hóa đơn", "Kết quả kiểm tra hóa đơn"]

Q_ROWS = [
    [1.0, "1", "C26TAA", "299", "22/09/2026", "0100112233", "CÔNG TY A", "1201496252", "CÔNG TY B",
     "Đồng Tháp", 3694444.0, 295556.0, 0.0, None, 3990000.0, "VND", "1.0", "Hóa đơn mới",
     "Đã cấp mã hóa đơn"],
    [2.0, "2", "C26TSA", "1545", "01/10/2026", "0104947250-001", "CHI NHÁNH C", "1201496252",
     "CÔNG TY B", "Đồng Tháp", None, None, 0.0, 234362.0, 237057.0, "VND", "1.0",
     "Hóa đơn thay thế", "Đã cấp mã hóa đơn"],
]
SCO_ROW = [1.0, "1", "C26MAA", "1237378", "03/08/2026", "0302249586", "SIÊU THỊ D", "Khu B, TP HCM",
           "1201496252", "CONG TY B", None, 1240924.0, 99274.0, 0.0, 1340198.0, "Hóa đơn mới",
           "Cục Thuế đã nhận hóa đơn có mã khởi tạo từ máy tính tiền"]


def _dimension_a1(data: bytes) -> bytes:
    """Ghi đè <dimension> thành A1 như file thật của GDT."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = re.sub(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1"/>', body)
            dst.writestr(item, body)
    return out.getvalue()


def gdt_xlsx(header, rows, title="DANH SÁCH HÓA ĐƠN") -> bytes:
    """Tiêu đề dòng 3, khoảng ngày dòng 4, tên cột dòng 6, dữ liệu từ dòng 7."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet 1"
    ws["A3"] = title
    ws["A4"] = "Từ ngày 21/09/2026 đến ngày 20/10/2026"
    for c, name in enumerate(header, 1):
        ws.cell(6, c, name)
    for r, values in enumerate(rows, 7):
        for c, v in enumerate(values, 1):
            if v is not None:
                ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return _dimension_a1(buf.getvalue())


XLSX = gdt_xlsx(QUERY_HEADER, Q_ROWS)

# ---------------------------------------------------------------------------
# Đọc file
# ---------------------------------------------------------------------------


def test_reads_every_row_despite_bogus_dimension():
    with zipfile.ZipFile(io.BytesIO(XLSX)) as z:                       # mẫu giống file GDT
        assert b'<dimension ref="A1"/>' in z.read("xl/worksheets/sheet1.xml")
    rows = invoice_list_excel.parse(XLSX)
    assert len(rows) == 2
    r = rows[0]
    assert (r["khmshdon"], r["khhdon"], r["shdon"]) == (1, "C26TAA", 299)
    assert r["nbmst"] == "0100112233" and r["nmmst"] == "1201496252" and r["nbten"] == "CÔNG TY A"
    assert r["tdlap"] == "2026-09-21T17:00:00Z"            # 0h 22/09 giờ VN, dạng UTC như API JSON
    assert r["tgtttbso"] == 3990000.0 and r["tgtphi"] is None and r["tgia"] == 1.0
    assert (r["tthai"], r["ttxly"]) == (1, 5)
    assert r["_tthai_text"] == "Hóa đơn mới" and r["_nguon"] == "excel"
    assert (rows[1]["khmshdon"], rows[1]["tthai"]) == (2, 2)


def test_excel_rows_feed_the_same_normalizer_as_json():
    r = api_client.parse_invoice_list_excel(XLSX, "mua")[1]
    n = job_policy.normalize_list_row(r, "1201496252", "mua")
    assert n["ngay_lap"] == date(2026, 10, 1)                 # ngày 1 không bị xếp sang tháng trước
    assert n["tong_tien"] == Decimal("237057.0") and n["tien_chua_thue"] is None
    assert (n["sub_type"], n["loai_hd"], n["direction"]) == ("mua", "Có mã", "mua")
    assert n["tthai"] == 2 and n["hthdon"] is None and not job_policy.summary_only(n)
    assert n["list_raw"]["_sub_type"] == "mua"


def test_reads_mtt_layout_by_column_name():
    rows = invoice_list_excel.parse(gdt_xlsx(SCO_HEADER, [SCO_ROW], "DANH SÁCH HÓA ĐƠN MÁY TÍNH TIỀN"))
    r = rows[0]
    assert (r["khhdon"], r["shdon"], r["nbmst"], r["nmmst"]) == (
        "C26MAA", 1237378, "0302249586", "1201496252")
    assert r["nbdchi"] == "Khu B, TP HCM" and "tgtphi" not in r and "dvtte" not in r
    assert r["tgtcthue"] == 1240924.0 and r["tgtttbso"] == 1340198.0
    assert r["ttxly"] == 8 and r["tdlap"] == "2026-08-02T17:00:00Z"


@pytest.mark.parametrize("text, code", [
    ("Hóa đơn mới", 1), ("Hóa đơn thay thế", 2), ("Hóa đơn điều chỉnh", 3),
    ("Hóa đơn đã bị thay thế", 4), ("Hóa đơn đã bị điều chỉnh", 5), ("Hóa đơn đã bị hủy", 6),
    ("Trạng thái lạ", None),
])
def test_status_text_to_code(text, code):
    row = list(Q_ROWS[0])
    row[17] = text
    r = invoice_list_excel.parse(gdt_xlsx(QUERY_HEADER, [row]))[0]
    assert r["tthai"] == code and r["_tthai_text"] == text


def test_check_result_text_to_ttxly():
    assert invoice_list_excel._ttxly("Đã cấp mã hóa đơn") == 5
    assert invoice_list_excel._ttxly("Tổng cục Thuế đã nhận không mã") == 6
    assert invoice_list_excel._ttxly("Cục Thuế đã nhận hóa đơn có mã khởi tạo từ máy tính tiền") == 8
    assert invoice_list_excel._ttxly(None) is None


def test_numeric_mst_keeps_leading_zero_and_new_columns_are_kept():
    row = list(Q_ROWS[0])
    row[5] = 100112233                                     # ô kiểu số → mất số 0 đầu
    row.append("1_C26TAA_298")
    r = invoice_list_excel.parse(gdt_xlsx(QUERY_HEADER + ["Hóa đơn liên quan"], [row]))[0]
    assert r["nbmst"] == "0100112233"
    assert r["Hóa đơn liên quan"] == "1_C26TAA_298"


def test_empty_list_is_not_an_error():
    assert invoice_list_excel.parse(gdt_xlsx(QUERY_HEADER, [])) == []


def test_wrong_files_are_rejected():
    with pytest.raises(ValueError):
        invoice_list_excel.parse(b'{"message":"loi"}')
    with pytest.raises(ValueError):                         # Excel khác, không có dòng tên cột
        invoice_list_excel.parse(gdt_xlsx(["A", "B"], [[1, 2]]))
    with pytest.raises(ValueError, match="ngày lập"):
        invoice_list_excel.parse(gdt_xlsx([h for h in QUERY_HEADER if h != "Ngày lập"], []))
    bad_date = list(Q_ROWS[0])
    bad_date[4] = "2026-09-22"
    with pytest.raises(ValueError, match="dd/MM/yyyy"):
        invoice_list_excel.parse(gdt_xlsx(QUERY_HEADER, [bad_date]))
    with pytest.raises(api_client.GDTError) as ei:
        api_client.parse_invoice_list_excel(b"PK-khong-phai-xlsx", "mua")
    assert ei.value.kind == "bad_content"


# ---------------------------------------------------------------------------
# Gọi GDT: URL giống cổng, phân loại lỗi
# ---------------------------------------------------------------------------

class _Resp:
    def __init__(self, status, content=b"", body=None):
        self.status_code = status
        self.content = content
        self.ok = 200 <= status < 300
        self.text = content.decode(errors="ignore")
        self._body = body or {}

    def json(self):
        return self._body


SCO_CAPTCHA = {"ckey": "k1", "cvalue": "AB12C"}


def _capture(monkeypatch, content=XLSX):
    seen = {}

    def fake_get(url, headers=None, params=None, timeout=None):
        seen.update(url=url, headers=headers, params=params)
        return _Resp(200, content)
    monkeypatch.setattr(api_client._gdt_session, "get", fake_get)
    monkeypatch.setattr(api_client, "sco_captcha", lambda: dict(SCO_CAPTCHA))   # sco-query bắt captcha (10/2026)
    return seen


# 3 URL chép từ trình duyệt khi bấm "Xuất Excel" trên cổng (20/10/2026)
PORTAL_URLS = {
    ("ban", "2026-09-21", "2026-10-20"):
        "https://hoadondientu.gdt.gov.vn/api/query/invoices/export-excel?sort=tdlap:desc"
        "&search=tdlap=ge=21/09/2026T00:00:00;tdlap=le=20/10/2026T23:59:59",
    ("mua", "2026-09-21", "2026-10-20"):
        "https://hoadondientu.gdt.gov.vn/api/query/invoices/export-excel-sold?sort=tdlap:desc"
        "&search=tdlap=ge=21/09/2026T00:00:00;tdlap=le=20/10/2026T23:59:59;ttxly==5%20%20%20%20"
        "&type=purchase",
    ("mua-mtt", "2026-08-01", "2026-08-31"):
        "https://hoadondientu.gdt.gov.vn/api/sco-query/invoices/export-excel-sold?sort=tdlap:desc"
        "&search=tdlap=ge=01/08/2026T00:00:00;tdlap=le=31/08/2026T23:59:59;ttxly==8%20%20%20%20"
        "&type=purchase&ckey=k1&cvalue=AB12C",       # từ 10/2026 cổng gửi kèm captcha (popup withCaptcha)
}


@pytest.mark.parametrize("key", list(PORTAL_URLS))
def test_export_request_equals_portal_url(monkeypatch, key):
    seen = _capture(monkeypatch)
    assert api_client.export_invoice_list("tok", *key) == XLSX
    ours = urlsplit(requests.Request("GET", seen["url"], params=seen["params"]).prepare().url)
    portal = urlsplit(PORTAL_URLS[key])
    assert ours.path == portal.path
    assert parse_qs(ours.query) == parse_qs(portal.query)     # so sau khi giải mã %20 / +
    assert seen["headers"]["Authorization"] == "Bearer tok" and seen["headers"]["request-id"]


@pytest.mark.parametrize("sub_type, path, ttxly", [
    ("mua-khma", "/query/invoices/export-excel-sold", "6"),
    ("ban-mtt", "/sco-query/invoices/export-excel", None),
])
def test_export_url_for_types_without_portal_sample(monkeypatch, sub_type, path, ttxly):
    seen = _capture(monkeypatch)
    api_client.export_invoice_list("tok", sub_type, "2026-09-21", "2026-10-20")
    assert seen["url"] == api_client.API_BASE + path
    assert seen["params"]["search"].endswith(f";ttxly=={ttxly}    " if ttxly else "T23:59:59")
    assert seen["params"].get("type") == ("purchase" if sub_type.startswith("mua") else None)
    assert ("ckey" in seen["params"]) == path.startswith("/sco-query")


@pytest.mark.parametrize("resp, kind", [
    (_Resp(401), "auth"),
    (_Resp(403), "auth"),
    (_Resp(429), "rate_limit"),
    (_Resp(503), "server"),
    (_Resp(500, b'{"message":"loi"}', {"message": "loi"}), "other"),
    (_Resp(404), "other"),
    (_Resp(200, b'{"message":"loi"}', {"message": "loi"}), "bad_content"),
])
def test_export_classifies_errors(monkeypatch, resp, kind):
    calls = []
    monkeypatch.setattr(api_client._gdt_session, "get", lambda *a, **k: calls.append(1) or resp)
    monkeypatch.setattr(api_client.time, "sleep", lambda s: None)
    with pytest.raises(api_client.GDTError) as ei:
        api_client.export_invoice_list("tok", "mua", "2026-09-21", "2026-10-20")
    assert ei.value.kind == kind
    assert len(calls) == (3 if ei.value.retryable else 1)    # chỉ lỗi tạm thời mới thử lại


def test_export_timeout_is_network_and_retried(monkeypatch):
    calls = []

    def boom(*a, **k):
        calls.append(1)
        raise requests.exceptions.Timeout()
    monkeypatch.setattr(api_client._gdt_session, "get", boom)
    monkeypatch.setattr(api_client.time, "sleep", lambda s: None)
    with pytest.raises(api_client.GDTError) as ei:
        api_client.export_invoice_list("tok", "ban", "2026-09-21", "2026-10-20")
    assert ei.value.kind == "network" and len(calls) == 3


# ---------------------------------------------------------------------------
# Job: lấy nguyên danh sách → hộp chờ
# ---------------------------------------------------------------------------

class _Tokens:
    mst = "1201496252"

    def get(self):
        return "tok"

    def refresh(self):
        return "tok"


START, END = date(2026, 9, 21), date(2026, 10, 20)


def _no_call(*a, **k):
    raise AssertionError("không được gọi")


def _raise(kind):
    def f(*a, **k):
        raise api_client.GDTError(f"lỗi {kind}", kind)
    return f


@pytest.fixture
def excel_source(monkeypatch, tmp_path):
    monkeypatch.delenv("INVOICE_LIST_SOURCE", raising=False)
    monkeypatch.setattr(invoice_job, "output_dir", lambda: tmp_path)
    return tmp_path


def test_list_step_saves_original_file_then_reads_it(monkeypatch, excel_source):
    monkeypatch.setattr(invoice_job.api_client, "export_invoice_list", lambda tok, st, f, t: XLSX)
    monkeypatch.setattr(invoice_job.api_client, "fetch_invoices_subtype", _no_call)
    stats = Counter()
    rows = invoice_job.fetch_list(_Tokens(), "1201496252", "mua", START, END, stats)
    assert [r["shdon"] for r in rows] == [299, 1545] and rows[0]["_sub_type"] == "mua"
    saved = excel_source / "list" / "1201496252" / "mua_20260921_20261020.xlsx"
    assert saved.read_bytes() == XLSX and not stats


@pytest.mark.parametrize("export, saved", [
    (_raise("other"), False),                                     # GDT từ chối xuất (404/500)
    (lambda *a: gdt_xlsx(["A", "B"], [[1, 2]]), True),            # file lạ: vẫn giữ để xem lại
])
def test_list_step_falls_back_to_json_when_export_unusable(monkeypatch, excel_source, export, saved):
    monkeypatch.setattr(invoice_job.api_client, "export_invoice_list", export)
    monkeypatch.setattr(invoice_job.api_client, "fetch_invoices_subtype",
                        lambda tok, st, f, t: [{"nbmst": "0100112233", "khhdon": "C26TAA", "shdon": 1}])
    stats = Counter()
    rows = invoice_job.fetch_list(_Tokens(), "1201496252", "mua", START, END, stats)
    assert rows[0]["shdon"] == 1 and stats["list_json_fallback"] == 1
    assert (excel_source / "list").exists() == saved


def test_list_step_does_not_hide_gdt_outage(monkeypatch, excel_source):
    """GDT quá tải: không gọi thêm API JSON (cũng sẽ lỗi) — để sync_lists dừng + ngắt sớm."""
    monkeypatch.setattr(invoice_job.api_client, "export_invoice_list", _raise("server"))
    monkeypatch.setattr(invoice_job.api_client, "fetch_invoices_subtype", _no_call)
    with pytest.raises(api_client.GDTError):
        invoice_job.fetch_list(_Tokens(), "1201496252", "mua", START, END, Counter())


def test_list_source_json_skips_excel(monkeypatch):
    monkeypatch.setenv("INVOICE_LIST_SOURCE", "json")
    monkeypatch.setattr(invoice_job.api_client, "export_invoice_list", _no_call)
    monkeypatch.setattr(invoice_job.api_client, "fetch_invoices_subtype", lambda tok, st, f, t: [])
    assert invoice_job.fetch_list(_Tokens(), "1201496252", "mua", START, END, Counter()) == []


class _BoxStore:
    def __init__(self):
        self.box, self.cursor = [], {}

    def upsert_invoices(self, rows):
        self.box.extend(rows)
        return len(rows)

    def set_cursor(self, mst, st, d):
        self.cursor[st] = d


class _DbLikeStore(_BoxStore):
    """upsert_invoices như SQL thật (ON CONFLICT DO NOTHING): HĐ đã có → bỏ qua, không xóa gì."""

    def __init__(self):
        super().__init__()
        self.keys = set()

    def upsert_invoices(self, rows):
        new = 0
        for n in rows:
            key = (n["nbmst"], n["khmshdon"], n["khhdon"], n["shdon"])
            if key not in self.keys:
                self.keys.add(key)
                self.box.append(n)
                new += 1
        return new


def test_list_differs_each_time_box_only_gets_new_invoices(monkeypatch, excel_source, caplog):
    def gdt_list(*shdons):
        return gdt_xlsx(QUERY_HEADER, [Q_ROWS[0][:3] + [str(s)] + Q_ROWS[0][4:] for s in shdons])
    fetches = iter([gdt_list(299, 300), gdt_list(300, 301, 301), gdt_list(299)])   # đủ / trùng / thiếu
    monkeypatch.setattr(invoice_job.api_client, "export_invoice_list", lambda tok, st, f, t: next(fetches))
    caplog.set_level("INFO", logger="invoice_job")
    store, stats = _DbLikeStore(), Counter()
    for _ in range(3):
        invoice_job.sync_lists(store, {"company_mst": "1201496252"}, _Tokens(), ["mua"],
                               lambda st: [(START, END)], job_policy.CircuitBreaker(), stats)
    assert [n["shdon"] for n in store.box] == [299, 300, 301]       # mỗi HĐ vào hộp đúng 1 lần
    assert stats["new_invoices"] == 3 and stats["listed"] == 6
    assert "GDT trả 2 HĐ (1 dòng trùng) → đã có 1 (bỏ qua), mới 1 (vào hộp chờ)" in caplog.text
    assert "GDT trả 1 HĐ → đã có 1 (bỏ qua), mới 0 (vào hộp chờ)" in caplog.text   # 300, 301 vắng: giữ


def test_whole_list_goes_into_box(monkeypatch, excel_source):
    files = {"ban": XLSX, "mua-mtt": gdt_xlsx(SCO_HEADER, [SCO_ROW])}
    monkeypatch.setattr(invoice_job.api_client, "export_invoice_list", lambda tok, st, f, t: files[st])
    store, stats = _BoxStore(), Counter()
    invoice_job.sync_lists(store, {"company_mst": "1201496252"}, _Tokens(), ["ban", "mua-mtt"],
                           lambda st: [(START, END)], job_policy.CircuitBreaker(), stats)
    assert [(n["sub_type"], n["khhdon"], n["shdon"]) for n in store.box] == [
        ("ban", "C26TAA", 299), ("ban", "C26TSA", 1545), ("mua-mtt", "C26MAA", 1237378)]
    assert store.box[0]["ngay_lap"] == date(2026, 9, 22) and store.box[0]["direction"] == "ban"
    assert store.cursor == {"ban": END, "mua-mtt": END}
    assert stats["listed"] == 3 and stats["new_invoices"] == 3


# ---------------------------------------------------------------------------
# Bước XML: HĐ không mã không có XML trên Thuế (danh sách Excel không có hthdon)
# ---------------------------------------------------------------------------

class _XmlQueue:
    def __init__(self, tasks):
        self.pending = list(tasks)
        self.parked, self.failed, self.released = [], [], []

    def claim_tasks(self, task_type, mst, n):
        batch, self.pending = self.pending[:n], self.pending[n:]
        return batch

    def mark_no_source(self, task, reason, run_id, kind="no_source"):
        self.parked.append((task["task_id"], kind))

    def fail_task(self, task, kind, error, run_id):
        self.failed.append((task["task_id"], kind))
        return "pending"

    def release_tasks(self, ids):
        self.released.extend(ids)


def test_k_invoice_without_xml_waits_for_mail_instead_of_retrying(monkeypatch, tmp_path):
    def no_xml(tok, row):
        raise api_client.GDTError("XML không có sẵn trên GDT (500)", kind="not_ready")
    monkeypatch.setattr(invoice_job.api_client, "download_invoice_zip", no_xml)
    monkeypatch.setattr(invoice_job, "XML_THROTTLE_SECONDS", 0)
    base = {"nbmst": "0300555450", "khmshdon": 1, "sub_type": "mua-khma", "direction": "mua",
            "ngay_lap": date(2026, 10, 9), "attempts": 0}
    tasks = [dict(base, task_id=1, khhdon="K26TAN", shdon=1, hthdon=None),   # từ Excel → chờ mail
             dict(base, task_id=2, khhdon="C26TAA", shdon=2, hthdon=None),   # có mã: XML có sau vài ngày
             dict(base, task_id=3, khhdon="K26TAN", shdon=3, hthdon=0)]      # API JSON: Thuế có XML
    store, stats = _XmlQueue(tasks), Counter()
    invoice_job.process_xml(store, {"company_mst": "1201496252"}, _Tokens(),
                            job_policy.CircuitBreaker(), 1, 10, stats, tmp_path)
    assert store.parked == [(1, "summary_only")]
    assert store.failed == [(2, "not_ready"), (3, "not_ready")]
    assert stats["xml_summary_only"] == 1 and stats["xml_failed"] == 2
